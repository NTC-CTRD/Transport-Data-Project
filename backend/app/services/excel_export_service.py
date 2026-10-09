from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.worksheet.table import Table, TableStyleInfo

from sqlalchemy.orm import Session

from app.models.corridor import Corridor
from app.models.observation import TrafficObservation
from app.models.segment_observation import SegmentTrafficObservation
from app.models.corridor_stop import CorridorStop


def _format_worksheet(worksheet):
    """
    Apply basic formatting to an Excel worksheet.
    """

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    worksheet.freeze_panes = "A2"

    for column_cells in worksheet.columns:
        max_length = 0

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value)),
                )

        column_letter = get_column_letter(
            column_cells[0].column
        )

        worksheet.column_dimensions[
            column_letter
        ].width = min(max_length + 2, 40)


def _format_datetime(value):
    """
    Convert timezone-aware datetime to an
    Excel-compatible datetime.
    """

    if value is None:
        return None

    if value.tzinfo is not None:
        return value.replace(tzinfo=None)

    return value


def _add_table(worksheet, table_name):
    """
    Add an Excel table to a worksheet when
    the worksheet contains data.
    """

    if worksheet.max_row < 2:
        return

    last_column = get_column_letter(
        worksheet.max_column
    )

    table_reference = (
        f"A1:{last_column}{worksheet.max_row}"
    )

    table = Table(
        displayName=table_name,
        ref=table_reference,
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    table.tableStyleInfo = style

    worksheet.add_table(table)


def _add_traffic_heat_map(
    workbook,
    db,
    corridor_id=None,
    start_date=None,
    end_date=None,
):
    """
    Create a traffic-condition heat map.

    Rows:
        Observation date/time

    Columns:
        Corridor segments

    Values:
        Traffic condition
    """

    sheet = workbook.create_sheet(
        "Traffic Heat Map"
    )

    segment_query = (
        db.query(
            SegmentTrafficObservation,
            Corridor,
            CorridorStop,
        )
        .join(
            Corridor,
            SegmentTrafficObservation.corridor_id
            == Corridor.id,
        )
        .join(
            CorridorStop,
            SegmentTrafficObservation.from_stop_id
            == CorridorStop.id,
        )
        .filter(
            Corridor.status == "Active"
        )
        .order_by(
            SegmentTrafficObservation.observed_at.asc(),
            SegmentTrafficObservation.sequence_order.asc(),
        )
    )

    if corridor_id is not None:
        segment_query = segment_query.filter(
            SegmentTrafficObservation.corridor_id
            == corridor_id
        )

    if start_date is not None:
        segment_query = segment_query.filter(
            SegmentTrafficObservation.observed_at
            >= start_date
        )

    if end_date is not None:
        segment_query = segment_query.filter(
            SegmentTrafficObservation.observed_at
            <= end_date
        )

    records = segment_query.all()

    if not records:
        sheet.append(
            [
                "No segment traffic data available."
            ]
        )
        _format_worksheet(sheet)
        return

    # This heat map is intended for one corridor.
    corridor = records[0][1]

    # Build segment names.
    segment_names = {}

    for segment, corridor_obj, from_stop in records:

        to_stop = (
            db.query(CorridorStop)
            .filter(
                CorridorStop.id
                == segment.to_stop_id
            )
            .first()
        )

        segment_names[
            segment.sequence_order
        ] = (
            f"{segment.sequence_order}. "
            f"{from_stop.name} → "
            f"{to_stop.name if to_stop else 'Unknown'}"
        )

    ordered_segments = sorted(
        segment_names.keys()
    )

    sheet.append(
        [
            "Date & Time"
        ]
        + [
            segment_names[number]
            for number in ordered_segments
        ]
    )

    # Group records by observation time.
    observations_by_time = {}

    for segment, corridor_obj, from_stop in records:

        observed_at = _format_datetime(
            segment.observed_at
        )

        if observed_at not in observations_by_time:
            observations_by_time[
                observed_at
            ] = {}

        observations_by_time[
            observed_at
        ][segment.sequence_order] = (
            segment.traffic_condition
        )

    for observed_at in sorted(
        observations_by_time.keys()
    ):

        row = [observed_at]

        for sequence_order in ordered_segments:
            row.append(
                observations_by_time[
                    observed_at
                ].get(sequence_order)
            )

        sheet.append(row)

    sheet.freeze_panes = "B2"

    # Date/time formatting.
    for cell in sheet["A"][1:]:
        cell.number_format = (
            "dd/mm/yyyy hh:mm"
        )

    # Traffic-condition colour rules.
    max_row = sheet.max_row
    max_column = sheet.max_column

    if max_row >= 2 and max_column >= 2:

        # Free Flow
        for row in sheet.iter_rows(
            min_row=2,
            min_col=2,
            max_row=max_row,
            max_col=max_column,
        ):
            for cell in row:

                if cell.value == "Free Flow":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="C6EFCE",
                    )

                elif cell.value == "Light":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="BDD7EE",
                    )

                elif cell.value == "Moderate":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="FFF2CC",
                    )

                elif cell.value == "Heavy":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="F4B183",
                    )

                elif cell.value == "Severe":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="F4CCCC",
                    )

    # Add title information.
    sheet.insert_rows(1)

    sheet["A1"] = (
        f"Traffic Heat Map — "
        f"{corridor.name}"
    )

    sheet["A1"].font = Font(
        bold=True,
        size=14,
    )

    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=max_column,
    )

    sheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    # Header is now row 2.
    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )
        cell.alignment = Alignment(
            horizontal="center"
        )

    # Column widths.
    sheet.column_dimensions["A"].width = 22

    for column in range(
        2,
        sheet.max_column + 1,
    ):
        sheet.column_dimensions[
            get_column_letter(column)
        ].width = 28


def _add_speed_heat_map(
    workbook,
    db,
    corridor_id=None,
    start_date=None,
    end_date=None,
):
    """
    Create an average-speed heat map.

    Rows:
        Observation date/time

    Columns:
        Corridor segments

    Values:
        Average speed in km/h
    """

    sheet = workbook.create_sheet(
        "Speed Heat Map"
    )

    segment_query = (
        db.query(
            SegmentTrafficObservation,
            Corridor,
            CorridorStop,
        )
        .join(
            Corridor,
            SegmentTrafficObservation.corridor_id
            == Corridor.id,
        )
        .join(
            CorridorStop,
            SegmentTrafficObservation.from_stop_id
            == CorridorStop.id,
        )
        .filter(
            Corridor.status == "Active"
        )
        .order_by(
            SegmentTrafficObservation.observed_at.asc(),
            SegmentTrafficObservation.sequence_order.asc(),
        )
    )

    if corridor_id is not None:
        segment_query = segment_query.filter(
            SegmentTrafficObservation.corridor_id
            == corridor_id
        )

    if start_date is not None:
        segment_query = segment_query.filter(
            SegmentTrafficObservation.observed_at
            >= start_date
        )

    if end_date is not None:
        segment_query = segment_query.filter(
            SegmentTrafficObservation.observed_at
            <= end_date
        )

    records = segment_query.all()

    if not records:
        sheet.append(
            [
                "No segment speed data available."
            ]
        )
        _format_worksheet(sheet)
        return

    corridor = records[0][1]

    segment_names = {}

    for segment, corridor_obj, from_stop in records:

        to_stop = (
            db.query(CorridorStop)
            .filter(
                CorridorStop.id
                == segment.to_stop_id
            )
            .first()
        )

        segment_names[
            segment.sequence_order
        ] = (
            f"{segment.sequence_order}. "
            f"{from_stop.name} → "
            f"{to_stop.name if to_stop else 'Unknown'}"
        )

    ordered_segments = sorted(
        segment_names.keys()
    )

    sheet.append(
        [
            "Date & Time"
        ]
        + [
            segment_names[number]
            for number in ordered_segments
        ]
    )

    observations_by_time = {}

    for segment, corridor_obj, from_stop in records:

        observed_at = _format_datetime(
            segment.observed_at
        )

        if observed_at not in observations_by_time:
            observations_by_time[
                observed_at
            ] = {}

        observations_by_time[
            observed_at
        ][segment.sequence_order] = (
            round(
                segment.average_speed_kmh,
                2,
            )
            if segment.average_speed_kmh
            is not None
            else None
        )

    for observed_at in sorted(
        observations_by_time.keys()
    ):

        row = [observed_at]

        for sequence_order in ordered_segments:
            row.append(
                observations_by_time[
                    observed_at
                ].get(sequence_order)
            )

        sheet.append(row)

    sheet.freeze_panes = "B2"

    for cell in sheet["A"][1:]:
        cell.number_format = (
            "dd/mm/yyyy hh:mm"
        )

    # Add a three-colour speed scale.
    if (
        sheet.max_row >= 2
        and sheet.max_column >= 2
    ):

        speed_range = (
            f"B2:"
            f"{get_column_letter(sheet.max_column)}"
            f"{sheet.max_row}"
        )

        color_scale = ColorScaleRule(
            start_type="min",
            start_color="F4CCCC",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFF2CC",
            end_type="max",
            end_color="C6EFCE",
        )

        sheet.conditional_formatting.add(
            speed_range,
            color_scale,
        )

    sheet.insert_rows(1)

    sheet["A1"] = (
        f"Speed Heat Map — "
        f"{corridor.name}"
    )

    sheet["A1"].font = Font(
        bold=True,
        size=14,
    )

    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=sheet.max_column,
    )

    sheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )
        cell.alignment = Alignment(
            horizontal="center"
        )

    sheet.column_dimensions["A"].width = 22

    for column in range(
        2,
        sheet.max_column + 1,
    ):
        sheet.column_dimensions[
            get_column_letter(column)
        ].width = 28


def generate_traffic_excel(
    db: Session,
    corridor_id: int | None = None,
    start_date: datetime | None = None,
    end_date: datetime | None = None,
):
    """
    Generate a complete M-TRADA Excel workbook.

    Sheets:
    1. Corridor Information
    2. Corridor Stops
    3. Traffic History
    4. Segment History
    5. Traffic Heat Map
    6. Speed Heat Map
    """

    workbook = Workbook()

    # ==========================================================
    # GET CORRIDORS
    # ==========================================================

    corridor_query = (
        db.query(Corridor)
        .filter(
            Corridor.status == "Active"
        )
        .order_by(
            Corridor.id.asc()
        )
    )

    if corridor_id is not None:
        corridor_query = corridor_query.filter(
            Corridor.id == corridor_id
        )

    corridors = corridor_query.all()

    # ==========================================================
    # SHEET 1 — CORRIDOR INFORMATION
    # ==========================================================

    corridor_sheet = workbook.active

    corridor_sheet.title = (
        "Corridor Information"
    )

    corridor_sheet.append(
        [
            "Corridor ID",
            "Corridor Name",
            "Origin",
            "Destination",
            "Distance (km)",
            "Status",
        ]
    )

    for corridor in corridors:

        corridor_sheet.append(
            [
                corridor.id,
                corridor.name,
                corridor.origin,
                corridor.destination,
                (
                    round(
                        corridor.distance_km,
                        3,
                    )
                    if corridor.distance_km
                    is not None
                    else None
                ),
                corridor.status,
            ]
        )

    _format_worksheet(
        corridor_sheet
    )

    _add_table(
        corridor_sheet,
        "CorridorInformationTable",
    )

    # ==========================================================
    # SHEET 2 — CORRIDOR STOPS
    # ==========================================================

    stops_sheet = workbook.create_sheet(
        "Corridor Stops"
    )

    stops_sheet.append(
        [
            "Corridor ID",
            "Corridor Name",
            "Stop Order",
            "Stop Name",
            "Latitude",
            "Longitude",
        ]
    )

    for corridor in corridors:

        stops = (
            db.query(CorridorStop)
            .filter(
                CorridorStop.corridor_id
                == corridor.id
            )
            .order_by(
                CorridorStop.sequence_order.asc()
            )
            .all()
        )

        for stop in stops:

            stops_sheet.append(
                [
                    corridor.id,
                    corridor.name,
                    stop.sequence_order,
                    stop.name,
                    (
                        round(
                            stop.latitude,
                            6,
                        )
                        if stop.latitude
                        is not None
                        else None
                    ),
                    (
                        round(
                            stop.longitude,
                            6,
                        )
                        if stop.longitude
                        is not None
                        else None
                    ),
                ]
            )

    _format_worksheet(
        stops_sheet
    )

    _add_table(
        stops_sheet,
        "CorridorStopsTable",
    )

    # ==========================================================
    # SHEET 3 — TRAFFIC HISTORY
    # ==========================================================

    traffic_sheet = workbook.create_sheet(
        "Traffic History"
    )

    traffic_sheet.append(
        [
            "Observation ID",
            "Corridor ID",
            "Corridor",
            "Origin",
            "Destination",
            "Distance (km)",
            "Travel Time (min)",
            "Average Speed (km/h)",
            "Traffic Condition",
            "Observed At",
            "Data Source",
        ]
    )

    traffic_query = (
        db.query(
            TrafficObservation,
            Corridor,
        )
        .join(
            Corridor,
            TrafficObservation.corridor_id
            == Corridor.id,
        )
        .filter(
            Corridor.status == "Active"
        )
        .order_by(
            TrafficObservation.observed_at.asc()
        )
    )

    if corridor_id is not None:
        traffic_query = traffic_query.filter(
            TrafficObservation.corridor_id
            == corridor_id
        )

    if start_date is not None:
        traffic_query = traffic_query.filter(
            TrafficObservation.observed_at
            >= start_date
        )

    if end_date is not None:
        traffic_query = traffic_query.filter(
            TrafficObservation.observed_at
            <= end_date
        )

    traffic_records = traffic_query.all()

    for observation, corridor in traffic_records:

        traffic_sheet.append(
            [
                observation.id,
                corridor.id,
                corridor.name,
                corridor.origin,
                corridor.destination,
                (
                    round(
                        observation.distance_km,
                        3,
                    )
                    if observation.distance_km
                    is not None
                    else None
                ),
                (
                    round(
                        observation.travel_time_minutes,
                        2,
                    )
                    if observation.travel_time_minutes
                    is not None
                    else None
                ),
                (
                    round(
                        observation.average_speed_kmh,
                        2,
                    )
                    if observation.average_speed_kmh
                    is not None
                    else None
                ),
                observation.traffic_condition,
                _format_datetime(
                    observation.observed_at
                ),
                observation.data_source,
            ]
        )

    _format_worksheet(
        traffic_sheet
    )

    _add_table(
        traffic_sheet,
        "TrafficHistoryTable",
    )

    # ==========================================================
    # SHEET 4 — SEGMENT HISTORY
    # ==========================================================

    segment_sheet = workbook.create_sheet(
        "Segment History"
    )

    segment_sheet.append(
        [
            "Observation ID",
            "Corridor ID",
            "Corridor",
            "Segment",
            "From Stop",
            "To Stop",
            "Distance (km)",
            "Travel Time (min)",
            "Average Speed (km/h)",
            "Traffic Condition",
            "Traffic Duration (sec)",
            "Static Duration (sec)",
            "Observed At",
            "Data Source",
        ]
    )

    segment_query = (
        db.query(
            SegmentTrafficObservation,
            Corridor,
            CorridorStop,
        )
        .join(
            Corridor,
            SegmentTrafficObservation.corridor_id
            == Corridor.id,
        )
        .join(
            CorridorStop,
            SegmentTrafficObservation.from_stop_id
            == CorridorStop.id,
        )
        .filter(
            Corridor.status == "Active"
        )
        .order_by(
            SegmentTrafficObservation.observed_at.asc(),
            SegmentTrafficObservation.sequence_order.asc(),
        )
    )

    if corridor_id is not None:
        segment_query = segment_query.filter(
            SegmentTrafficObservation.corridor_id
            == corridor_id
        )

    if start_date is not None:
        segment_query = segment_query.filter(
            SegmentTrafficObservation.observed_at
            >= start_date
        )

    if end_date is not None:
        segment_query = segment_query.filter(
            SegmentTrafficObservation.observed_at
            <= end_date
        )

    segment_records = segment_query.all()

    for segment, corridor, from_stop in segment_records:

        to_stop = (
            db.query(CorridorStop)
            .filter(
                CorridorStop.id
                == segment.to_stop_id
            )
            .first()
        )

        segment_sheet.append(
            [
                segment.id,
                corridor.id,
                corridor.name,
                segment.sequence_order,
                from_stop.name,
                (
                    to_stop.name
                    if to_stop
                    else None
                ),
                (
                    round(
                        segment.distance_km,
                        3,
                    )
                    if segment.distance_km
                    is not None
                    else None
                ),
                (
                    round(
                        segment.travel_time_minutes,
                        2,
                    )
                    if segment.travel_time_minutes
                    is not None
                    else None
                ),
                (
                    round(
                        segment.average_speed_kmh,
                        2,
                    )
                    if segment.average_speed_kmh
                    is not None
                    else None
                ),
                segment.traffic_condition,
                (
                    round(
                        segment.traffic_duration_seconds,
                        2,
                    )
                    if segment.traffic_duration_seconds
                    is not None
                    else None
                ),
                (
                    round(
                        segment.static_duration_seconds,
                        2,
                    )
                    if segment.static_duration_seconds
                    is not None
                    else None
                ),
                _format_datetime(
                    segment.observed_at
                ),
                segment.data_source,
            ]
        )

    _format_worksheet(
        segment_sheet
    )

    _add_table(
        segment_sheet,
        "SegmentHistoryTable",
    )

    # ==========================================================
    # SHEET 5 — TRAFFIC HEAT MAP
    # ==========================================================

    _add_traffic_heat_map(
        workbook,
        db,
        corridor_id,
        start_date,
        end_date,
    )

    # ==========================================================
    # SHEET 6 — SPEED HEAT MAP
    # ==========================================================

    _add_speed_heat_map(
        workbook,
        db,
        corridor_id,
        start_date,
        end_date,
    )

    # ==========================================================
    # SAVE WORKBOOK TO MEMORY
    # ==========================================================

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output